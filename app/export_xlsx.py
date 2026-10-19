"""Readable Excel export of all listings for the admin.

One row per listing, sorted by employer and town, with columns grouped under coloured headings,
a sheet with the full ad texts and a sheet explaining every column. The CSV export stays the
machine-readable format for re-import.
"""
from __future__ import annotations

import io
import re
from collections.abc import Iterable
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Opportunity

NONE = "not stated"
EXPIRES_SOON_DAYS = 14

STATUS_LABELS = {
    "draft": "Draft", "submitted": "Submitted by employer", "awaiting_review": "Waiting for your review",
    "published": "Published", "expired": "Expired", "closed": "Closed", "rejected": "Rejected",
}
VERIFICATION_LABELS = {
    "unverified": "Not checked yet", "employer_confirmed": "Confirmed by employer",
    "official_source_verified": "Checked against official source", "admin_verified": "Checked by admin",
    "student_reported": "Reported by a student",
}

GROUPS = [
    ("LISTING", "1F4E79", [("No.", 6), ("Employer", 26), ("Internship title", 34)]),
    ("WHERE", "2E7D32", [("Postcode", 10), ("Town", 22)]),
    ("WHO CAN APPLY", "8E24AA", [("Age", 12), ("School class", 15)]),
    ("WHEN", "EF6C00", [("Dates", 22), ("Apply by", 12)]),
    ("HOW TO APPLY", "00838F", [("Documents to send", 30), ("Email", 22), ("Original ad", 14)]),
    ("FIELD", "6D4C41", [("Field", 30)]),
    ("REVIEW", "C62828", [("Status", 20), ("Checked?", 16), ("Shown until", 12), ("Things to check", 40)]),
]
LINK_COL, EXPIRES_COL, CHECKS_COL = 12, 16, 17

GUIDE = [
    ("LISTING", "No.", "The listing's number in the app. Search for it in the admin area to open it."),
    ("LISTING", "Employer", "The company offering the internship."),
    ("LISTING", "Internship title", "The title of the ad as the employer wrote it."),
    ("WHERE", "Postcode / Town", "Where the internship takes place."),
    ("WHO CAN APPLY", "Age", "Minimum and maximum age, if known. Grey 'not stated' means it isn't filled in."),
    ("WHO CAN APPLY", "School class", "Which school classes (Klasse) can apply, if known."),
    ("WHEN", "Dates", "Fixed internship periods, or 'Flexible' when students can agree dates with the employer."),
    ("WHEN", "Apply by", "Application deadline, if known."),
    ("HOW TO APPLY", "Documents to send", "Documents students should send (cover letter, CV, report card, school form...). "
                                          "For imported ads these were read from the ad text; check them against the ad."),
    ("HOW TO APPLY", "Email", "Application email address, if known."),
    ("HOW TO APPLY", "Original ad", "Click to open the employer's own ad."),
    ("FIELD", "Field", "The career fields of the listing. For imported ads the app guesses them from keywords, "
                       "which is often wrong; correct them when you review."),
    ("REVIEW", "Status", "'Waiting for your review' means students can't see it yet. Only 'Published' listings are visible."),
    ("REVIEW", "Checked?", "Whether someone has verified the listing. 'Not checked yet' for everything just imported."),
    ("REVIEW", "Shown until", "The date the listing expires automatically. Red means this date has already passed."),
    ("REVIEW", "Things to check", "Automatic hints for your review. Yellow means there is something to look at."),
    ("", "", ""),
    ("SHEETS", "Listings", "One row per listing, sorted by employer and then town. Rows of the same employer share a "
                          "background colour. Use the arrows in the header row to filter, for example by status."),
    ("SHEETS", "Full ad text", "The complete ad text for each listing, by number."),
]

THIN = Side(style="thin", color="BFBFBF")
THICK = Side(style="medium", color="404040")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
GREY = Font(color="9E9E9E", italic=True)
LINK = Font(color="0563C1", underline="single")
RED = PatternFill("solid", fgColor="FFCDD2")
AMBER = PatternFill("solid", fgColor="FFF3C4")
BAND = PatternFill("solid", fgColor="F3F6FA")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def short_employer(name: str) -> str:
    """'dm-drogerie markt GmbH + Co. KG' -> 'dm-drogerie markt GmbH'; other names stay as they are."""
    return re.sub(r"\s*\+\s*Co\.?\s*KG.*$", "", name or "").strip()


def age_text(lo: int | None, hi: int | None) -> str:
    if lo and hi:
        return f"{lo} years" if lo == hi else f"{lo} to {hi} years"
    if lo:
        return f"{lo} or older"
    if hi:
        return f"up to {hi} years"
    return NONE


def grade_text(lo: int | None, hi: int | None) -> str:
    if lo and hi:
        return f"Class {lo}" if lo == hi else f"Class {lo} to {hi}"
    if lo:
        return f"Class {lo} and up"
    if hi:
        return f"Up to class {hi}"
    return NONE


def dates_text(o: Opportunity) -> str:
    parts = [f"{w.start_date:%d.%m.%Y} to {w.end_date:%d.%m.%Y}" for w in sorted(o.date_windows, key=lambda w: w.start_date)]
    if parts:
        return "\n".join(parts)
    return "Flexible, agree with employer" if o.flexible_dates else NONE


def _day(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else NONE


def plain_text(text: str | None) -> str:
    text = re.sub(r"<[^>]+>", " ", text or "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def review_hints(o: Opportunity, today: date) -> list[str]:
    out = []
    if o.expires_at and o.expires_at < today:
        out.append("Expiry date already passed: ad is probably old")
    elif o.expires_at and (o.expires_at - today).days <= EXPIRES_SOON_DAYS:
        out.append("Expires within 2 weeks")
    if o.application_deadline and o.application_deadline < today:
        out.append("Application deadline already passed")
    if o.link_status == "broken":
        out.append("Link to the ad is broken")
    if not (o.min_age or o.max_age or o.min_grade or o.max_grade):
        out.append("Age and class not stated: check the ad")
    text = (o.description or "").lower()
    if "studi" in text or "hochschul" in text:
        out.append("Text also mentions university students: check it fits school students")
    return out


def _header(ws, row: int, values: list[tuple[str, int]], color: str = "1F4E79") -> None:
    for j, (name, width) in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=name)
        c.font = WHITE_BOLD
        c.fill = _fill(color)
        ws.column_dimensions[get_column_letter(j)].width = width


def _listings_sheet(ws, opps: list[Opportunity], today: date) -> None:
    col = 1
    group_edges = set()
    for title, color, cols in GROUPS:
        first = col
        for name, width in cols:
            c = ws.cell(row=2, column=col, value=name)
            c.font = WHITE_BOLD
            c.fill = _fill(color)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width
            col += 1
        ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=col - 1)
        g = ws.cell(row=1, column=first, value=title)
        g.font = Font(bold=True, color="FFFFFF", size=12)
        g.fill = _fill(color)
        g.alignment = Alignment(horizontal="center", vertical="center")
        group_edges.add(col - 1)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 32

    prev, shade = None, False
    for i, o in enumerate(opps, start=3):
        employer = short_employer(o.employer.name)
        if employer != prev:
            shade, prev = not shade, employer
        hints = review_hints(o, today)
        values = [
            o.id, employer, o.title, o.postcode_code or NONE, o.city_name or o.location_label or NONE,
            age_text(o.min_age, o.max_age), grade_text(o.min_grade, o.max_grade),
            dates_text(o), _day(o.application_deadline),
            "\n".join(d.name for d in o.documents) or NONE,
            o.application_email or NONE,
            "Open ad" if o.official_url else NONE,
            ", ".join(d.name for d in o.domains) or NONE,
            STATUS_LABELS.get(o.status, o.status),
            VERIFICATION_LABELS.get(o.verification_status, o.verification_status),
            _day(o.expires_at),
            "\n".join(hints) or "Nothing unusual",
        ]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = TOP_WRAP
            c.border = Border(left=THIN, right=THICK if j in group_edges else THIN, top=THIN, bottom=THIN)
            if shade:
                c.fill = BAND
            if v == NONE:
                c.font = GREY
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="center", vertical="top")
        if o.official_url:
            link = ws.cell(row=i, column=LINK_COL)
            link.hyperlink = o.official_url
            link.font = LINK
        if o.expires_at and o.expires_at < today:
            ws.cell(row=i, column=EXPIRES_COL).fill = RED
            ws.cell(row=i, column=CHECKS_COL).fill = RED
        elif hints:
            ws.cell(row=i, column=CHECKS_COL).fill = AMBER

    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(col - 1)}{max(len(opps), 1) + 2}"
    ws.sheet_view.zoomScale = 90


def build_workbook(opportunities: Iterable[Opportunity], today: date | None = None) -> bytes:
    today = today or date.today()
    opps = sorted(opportunities, key=lambda o: (short_employer(o.employer.name).lower(),
                                                (o.city_name or "").lower(), o.postcode_code or "", o.id))
    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"
    _listings_sheet(ws, opps, today)

    ds = wb.create_sheet("Full ad text")
    _header(ds, 1, [("No.", 6), ("Employer", 26), ("Town", 20), ("Ad text", 120)])
    for i, o in enumerate(opps, start=2):
        for j, v in enumerate([o.id, short_employer(o.employer.name), o.city_name or "", plain_text(o.description)], start=1):
            c = ds.cell(row=i, column=j, value=v)
            c.alignment = TOP_WRAP
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ds.freeze_panes = "B2"

    gs = wb.create_sheet("What each column means")
    _header(gs, 1, [("Group", 16), ("Column", 26), ("What it means", 100)])
    colors = {title: color for title, color, _ in GROUPS}
    for i, row in enumerate(GUIDE, start=2):
        for j, v in enumerate(row, start=1):
            gs.cell(row=i, column=j, value=v).alignment = TOP_WRAP
        if row[0] in colors:
            gs.cell(row=i, column=1).font = Font(bold=True, color=colors[row[0]])
        gs.cell(row=i, column=2).font = Font(bold=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
